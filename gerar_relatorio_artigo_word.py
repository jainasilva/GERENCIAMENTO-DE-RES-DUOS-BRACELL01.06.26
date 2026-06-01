from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


ROOT_DIR = Path(__file__).resolve().parent
PACOTE_DIR = ROOT_DIR / "PACOTE_STREAMLIT_CLOUD_BRACELL_ATUALIZADO"
SLIDES_FILE = PACOTE_DIR / "data" / "slides.json"
IMAGES_DIR = PACOTE_DIR / "assets" / "images"
OUTPUT_FILE = ROOT_DIR / "Relatorio_Artigo_Completo_BRACELL.docx"


def limpar_texto(texto: str) -> str:
    texto_limpo = str(texto or "").replace("\x0b", " ").strip()
    texto_limpo = re.sub(r"\s+", " ", texto_limpo).strip()
    return texto_limpo


def unicos_preservando_ordem(itens: list[str]) -> list[str]:
    vistos: set[str] = set()
    saida: list[str] = []
    for item in itens:
        chave = item.casefold()
        if not item or chave in vistos:
            continue
        vistos.add(chave)
        saida.append(item)
    return saida


def carregar_slides() -> list[dict]:
    dados = json.loads(SLIDES_FILE.read_text(encoding="utf-8"))
    slides: list[dict] = []
    for item in dados:
        if not isinstance(item, dict):
            continue
        numero = int(item.get("slide", 0) or 0)
        titulo = limpar_texto(item.get("title", ""))
        textos = [limpar_texto(t) for t in item.get("texts", []) if limpar_texto(t)]
        imagens = [limpar_texto(i) for i in item.get("images", []) if limpar_texto(i)]
        slides.append(
            {
                "slide": numero,
                "title": titulo,
                "texts": textos,
                "images": imagens,
            }
        )
    return sorted(slides, key=lambda s: s["slide"])


def encontrar_imagem_slide(slide: dict) -> Path | None:
    candidatos: list[Path] = []
    numero = int(slide.get("slide", 0) or 0)
    for ref in slide.get("images", []):
        candidatos.append(IMAGES_DIR / ref)
    if numero > 0:
        candidatos.extend(
            [
                IMAGES_DIR / f"slide{numero:02d}.png",
                IMAGES_DIR / f"slide{numero}.png",
            ]
        )
    for caminho in candidatos:
        if caminho.exists():
            return caminho
    return None


def configurar_estilo(documento: Document) -> None:
    estilo_normal = documento.styles["Normal"]
    estilo_normal.font.name = "Times New Roman"
    estilo_normal.font.size = Pt(12)
    for secao in documento.sections:
        secao.top_margin = Cm(3)
        secao.bottom_margin = Cm(2)
        secao.left_margin = Cm(3)
        secao.right_margin = Cm(2)


def adicionar_capa(documento: Document, slides: list[dict]) -> None:
    slide1 = slides[0] if slides else {"texts": [], "title": "Relatório Técnico"}
    linhas = [limpar_texto(t) for t in slide1.get("texts", [])]

    integrantes: list[str] = []
    curso = "ENGENHARIA CIVIL | UNISAGRADO"
    disciplina = "CIÊNCIAS E GESTÃO AMBIENTAL"
    for idx, linha in enumerate(linhas):
        if "INTEGRANTES DO GRUPO" in linha.upper():
            if idx + 1 < len(linhas):
                integrantes.append(linhas[idx + 1])
            if idx + 2 < len(linhas):
                integrantes.append(linhas[idx + 2])
        if linha.upper().startswith("CURSO:"):
            curso = linha.replace("CURSO:", "").strip()
        if linha.upper().startswith("DISCIPLINA:"):
            disciplina = linha.replace("DISCIPLINA:", "").strip()

    p_titulo = documento.add_paragraph("GESTÃO AMBIENTAL NA BRACELL")
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_titulo.runs[0].bold = True
    p_titulo.runs[0].font.size = Pt(18)

    p_sub = documento.add_paragraph(
        "Relatório em formato de artigo técnico com base no conteúdo completo da apresentação"
    )
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.runs[0].italic = True

    documento.add_paragraph("")
    p_inst = documento.add_paragraph("UNISAGRADO")
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.runs[0].bold = True

    p_curso = documento.add_paragraph(curso)
    p_curso.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_disc = documento.add_paragraph(f"Disciplina: {disciplina}")
    p_disc.alignment = WD_ALIGN_PARAGRAPH.CENTER

    documento.add_paragraph("")
    p_aut = documento.add_paragraph("Integrantes:")
    p_aut.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aut.runs[0].bold = True

    for grupo in unicos_preservando_ordem(integrantes):
        p_nome = documento.add_paragraph(grupo)
        p_nome.alignment = WD_ALIGN_PARAGRAPH.CENTER

    documento.add_paragraph("")
    p_data = documento.add_paragraph(
        f"Lençóis Paulista - SP, {date.today().strftime('%d/%m/%Y')}"
    )
    p_data.alignment = WD_ALIGN_PARAGRAPH.CENTER

    documento.add_page_break()


def adicionar_resumo(documento: Document) -> None:
    documento.add_heading("Resumo", level=1)
    documento.add_paragraph(
        "Este artigo apresenta uma análise integrada da gestão ambiental da Bracell, "
        "com foco em resíduos industriais, recursos hídricos, energia renovável, manejo florestal, "
        "licenciamento, passivos ambientais e programas de mitigação. O conteúdo foi organizado a partir "
        "da apresentação técnica do grupo, consolidando práticas operacionais, indicadores de desempenho "
        "e ações corretivas para prevenção de impactos. O estudo demonstra que a gestão ambiental estruturada "
        "fortalece a conformidade legal, melhora a eficiência de processos e amplia resultados socioambientais."
    )
    documento.add_paragraph(
        "Palavras-chave: gestão ambiental; resíduos industriais; recursos hídricos; passivo ambiental; Bracell."
    )


def adicionar_texto_slide(documento: Document, slide: dict) -> None:
    titulo = slide.get("title", f"Slide {slide.get('slide', '')}")
    texto_slide = slide.get("texts", [])
    linhas = [t for t in texto_slide if limpar_texto(t).casefold() != limpar_texto(titulo).casefold()]
    linhas = unicos_preservando_ordem(linhas)

    for linha in linhas:
        if linha.startswith("•"):
            documento.add_paragraph(linha.lstrip("• ").strip(), style="List Bullet")
            continue
        if re.match(r"^\-\s*", linha):
            documento.add_paragraph(re.sub(r"^\-\s*", "", linha), style="List Bullet")
            continue
        p = documento.add_paragraph(linha)
        if linha.lower().startswith("fonte:"):
            for run in p.runs:
                run.italic = True


def adicionar_corpo_principal(documento: Document, slides: list[dict]) -> None:
    mapa = {
        "1. Introdução e Objetivos": [2],
        "2. Política Ambiental e Caracterização da Empresa": [3, 4],
        "3. Gestão de Resíduos Industriais": [5, 6],
        "4. Recursos Hídricos e Energia": [7, 8, 9, 10, 11, 12],
        "5. Manejo, Licenciamento e Responsabilidade Socioambiental": [13, 14, 15, 16],
        "6. Passivos Ambientais e Planos de Ação": [17, 18, 19, 20, 21],
        "7. Indicadores de Sustentabilidade e Conformidade": [22, 23, 24, 25, 26, 27, 28, 29, 30],
        "8. Conclusão": [31],
    }
    por_numero = {int(s["slide"]): s for s in slides}

    for secao, numeros in mapa.items():
        documento.add_heading(secao, level=1)
        for numero in numeros:
            slide = por_numero.get(numero)
            if not slide:
                continue
            documento.add_heading(slide.get("title", f"Slide {numero}"), level=2)
            adicionar_texto_slide(documento, slide)


def extrair_referencias(slides: list[dict]) -> list[str]:
    por_numero = {int(s["slide"]): s for s in slides}
    refs: list[str] = []

    # Referências formais do último slide.
    slide_ref = por_numero.get(32, {})
    for linha in slide_ref.get("texts", []):
        texto = limpar_texto(linha).lstrip("• ").strip()
        if texto and texto.casefold() != limpar_texto(slide_ref.get("title", "")).casefold():
            refs.append(texto)

    # Complemento de fontes citadas nos demais slides.
    for slide in slides:
        linhas = slide.get("texts", [])
        for idx, linha in enumerate(linhas):
            if limpar_texto(linha).lower().startswith("fonte:"):
                atual = limpar_texto(linha)
                if idx + 1 < len(linhas):
                    prox = limpar_texto(linhas[idx + 1])
                    if prox and not prox.lower().startswith(("•", "-", "fonte:")):
                        atual = f"{atual} {prox}"
                refs.append(atual)

    refs = [r.replace("Fonte: Fonte:", "Fonte:") for r in refs]
    return unicos_preservando_ordem([limpar_texto(r) for r in refs if limpar_texto(r)])


def adicionar_referencias(documento: Document, slides: list[dict]) -> None:
    documento.add_heading("9. Referências", level=1)
    for ref in extrair_referencias(slides):
        texto = ref.lstrip("• ").strip()
        if texto:
            documento.add_paragraph(texto, style="List Number")


def adicionar_anexo(documento: Document, slides: list[dict]) -> None:
    documento.add_page_break()
    documento.add_heading("Anexo A - Conteúdo Integral da Apresentação", level=1)
    documento.add_paragraph(
        "Transcrição organizada slide a slide com inclusão das imagens de apoio da apresentação original."
    )

    for slide in slides:
        numero = int(slide.get("slide", 0) or 0)
        titulo = slide.get("title", f"Slide {numero}")
        documento.add_heading(f"Slide {numero}: {titulo}", level=2)
        adicionar_texto_slide(documento, slide)

        imagem = encontrar_imagem_slide(slide)
        if imagem:
            documento.add_picture(str(imagem), width=Cm(15.5))
            legenda = documento.add_paragraph(f"Figura {numero} - Reprodução do slide {numero}.")
            legenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in legenda.runs:
                run.italic = True


def main() -> None:
    if not SLIDES_FILE.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {SLIDES_FILE}")

    slides = carregar_slides()
    if not slides:
        raise ValueError("Nenhum slide foi carregado para montar o relatório.")

    doc = Document()
    configurar_estilo(doc)
    adicionar_capa(doc, slides)
    adicionar_resumo(doc)
    adicionar_corpo_principal(doc, slides)
    adicionar_referencias(doc, slides)
    adicionar_anexo(doc, slides)
    doc.save(OUTPUT_FILE)
    print(f"Relatório gerado com sucesso: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
